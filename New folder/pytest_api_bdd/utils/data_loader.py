"""
data_loader.py
==============
Single place to load test data from any external source.

Supported sources:
  - JSON  file   → load_json("booking_data.json", key="valid_booking")
  - Excel file   → load_excel("bookings.xlsx", sheet="CreateBooking", row=0)
  - .properties  → load_properties("config.properties", key="base_url")
  - BDD DataTable→ parse_datatable(table, as_format="list_of_dicts")

All paths are relative to  resource/test_data/
Tests never touch file I/O directly — they just call these functions.
"""

import json
import os
from pathlib import Path

# openpyxl is needed only for Excel — soft-import so JSON/properties still work
try:
    import openpyxl
    _EXCEL_AVAILABLE = True
except ImportError:
    _EXCEL_AVAILABLE = False

# Every data file lives here — one source of truth for paths
DATA_DIR = Path(__file__).resolve().parent.parent / "resource" / "test_data"


# ─────────────────────────────────────────────────────────────────────────────
#  JSON loader
# ─────────────────────────────────────────────────────────────────────────────

def load_json(filename: str, key: str = None) -> dict | list:
    """
    Load a JSON file from resource/test_data/.

    Args:
        filename : e.g. "booking_data.json"
        key      : optional top-level key to return, e.g. "valid_booking"
                   if None — returns the full file content

    Usage in test:
        payload = load_json("booking_data.json", key="valid_booking")
        response = api_client.post("/booking", payload=payload)
    """
    path = DATA_DIR / filename
    if not path.exists():
        raise FileNotFoundError(f"JSON data file not found: {path}")

    with open(path, encoding="utf-8") as f:
        data = json.load(f)

    if key:
        if key not in data:
            raise KeyError(f"Key '{key}' not found in {filename}. Available keys: {list(data.keys())}")
        return data[key]

    return data


# ─────────────────────────────────────────────────────────────────────────────
#  Excel loader
# ─────────────────────────────────────────────────────────────────────────────

def load_excel(filename: str, sheet: str = None, row: int = 0) -> dict:
    """
    Load ONE row from an Excel sheet as a dict.
    Row 0 = first data row (row index 1 in Excel, after header row 0).

    Args:
        filename : e.g. "bookings.xlsx"
        sheet    : sheet/tab name, e.g. "CreateBooking" — defaults to first sheet
        row      : 0-based index of the data row to return (not counting header)

    Usage in test:
        payload = load_excel("bookings.xlsx", sheet="CreateBooking", row=0)
        response = api_client.post("/booking", payload=payload)

    Excel format expected:
        | firstname | lastname | totalprice | depositpaid | checkin    | checkout   |
        | Yashwanth | Gowda    | 200        | TRUE        | 2025-01-01 | 2025-01-10 |

    Returns dict:
        {"firstname": "Yashwanth", "lastname": "Gowda", "totalprice": 200, ...}
    """
    if not _EXCEL_AVAILABLE:
        raise ImportError("openpyxl not installed. Run: pip install openpyxl")

    path = DATA_DIR / filename
    if not path.exists():
        raise FileNotFoundError(f"Excel file not found: {path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.active

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        raise ValueError(f"Sheet '{sheet}' has no data rows (only header or empty)")

    header = [str(h).strip() for h in rows[0]]  # first row = column names
    data_rows = rows[1:]                          # rest = data

    if row >= len(data_rows):
        raise IndexError(f"Row {row} out of range — sheet has {len(data_rows)} data row(s)")

    return dict(zip(header, data_rows[row]))


def load_excel_all_rows(filename: str, sheet: str = None) -> list[dict]:
    """
    Load ALL data rows from a sheet as a list of dicts.

    Usage with pytest.mark.parametrize:
        @pytest.mark.parametrize("payload", load_excel_all_rows("bookings.xlsx", sheet="CreateBooking"))
        def test_create_booking(api_client, payload):
            response = api_client.post("/booking", payload=payload)
            assert response.status_code == 200
    """
    if not _EXCEL_AVAILABLE:
        raise ImportError("openpyxl not installed. Run: pip install openpyxl")

    path = DATA_DIR / filename
    if not path.exists():
        raise FileNotFoundError(f"Excel file not found: {path}")

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.active

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return []

    header = [str(h).strip() for h in rows[0]]
    return [dict(zip(header, r)) for r in rows[1:]]


# ─────────────────────────────────────────────────────────────────────────────
#  .properties loader
# ─────────────────────────────────────────────────────────────────────────────

def load_properties(filename: str, key: str = None) -> str | dict:
    """
    Load a Java-style .properties file.

    Format:
        base_url=https://restful-booker.herokuapp.com
        username=admin
        password=password123

    Args:
        filename : e.g. "config.properties"
        key      : optional — return a single value by key
                   if None — returns ALL key/value pairs as a dict

    Usage:
        base_url = load_properties("config.properties", key="base_url")
        all_cfg  = load_properties("config.properties")
    """
    path = DATA_DIR / filename
    if not path.exists():
        raise FileNotFoundError(f".properties file not found: {path}")

    props = {}
    with open(path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#"):   # skip blanks and comments
                continue
            if "=" in line:
                k, _, v = line.partition("=")
                props[k.strip()] = v.strip()

    if key:
        if key not in props:
            raise KeyError(f"Key '{key}' not found in {filename}. Available: {list(props.keys())}")
        return props[key]

    return props


# ─────────────────────────────────────────────────────────────────────────────
#  BDD DataTable parser
#
#  Converts the raw pipe-delimited string from a feature file into
#  whichever Python format the step needs.
#
#  as_format options:
#    "list_of_dicts"  (default) — same as Java dataTable.asMaps()
#    "list_of_lists"            — same as Java dataTable.asLists()
#    "single_dict"              — key/value table (2 columns: key | value)
#    "flat_list"                — single column table, returns plain list
#
#  Feature file table (raw string passed by pytest-bdd):
#    | username | password    |
#    |          | password123 |
#    | admin    |             |
#    |          |             |
# ─────────────────────────────────────────────────────────────────────────────

def parse_datatable(table: str, as_format: str = "list_of_dicts"):
    """
    Parse a BDD pipe-delimited DataTable string.

    Args:
        table     : raw string from feature file
        as_format : one of —
                    "list_of_dicts"  → [{col1:val, col2:val}, ...]   (default)
                    "list_of_lists"  → [[val, val], [val, val], ...]
                    "single_dict"    → {key: value, key: value, ...}
                    "flat_list"      → [val, val, val, ...]

    Usage in step:
        rows = parse_datatable(table)                            # list of dicts
        rows = parse_datatable(table, as_format="list_of_lists") # list of lists
        cfg  = parse_datatable(table, as_format="single_dict")   # key/value dict
        vals = parse_datatable(table, as_format="flat_list")     # plain list
    """
    # ── parse raw string into 2D list ────────────────────────────────────────
    lines = [line.strip() for line in table.strip().splitlines() if line.strip()]
    grid  = [
        [cell.strip() for cell in line.strip("|").split("|")]
        for line in lines
    ]
    # grid example:
    # [["username", "password"],
    #  ["",         "password123"],
    #  ["admin",    ""],
    #  ["",         ""]]

    # ── convert to requested format ──────────────────────────────────────────

    if as_format == "list_of_dicts":
        # Java equivalent: dataTable.asMaps()
        # → [{"username":"", "password":"password123"}, ...]
        headers   = grid[0]
        data_rows = grid[1:]
        return [dict(zip(headers, row)) for row in data_rows]

    elif as_format == "list_of_lists":
        # Java equivalent: dataTable.asLists()
        # → [["username","password"], ["","password123"], ["admin",""], ...]
        # includes header row — caller slices [1:] if they want to skip it
        return grid

    elif as_format == "single_dict":
        # Use when table has exactly 2 columns: | key | value |
        # Java equivalent: dataTable.asMap()
        # Feature table:
        #   | firstname | Yashwanth |
        #   | lastname  | Gowda     |
        #   | price     | 200       |
        # → {"firstname":"Yashwanth", "lastname":"Gowda", "price":"200"}
        return {row[0]: row[1] for row in grid if len(row) == 2}

    elif as_format == "flat_list":
        # Use when table has a single column (no header needed)
        # Feature table:
        #   | admin        |
        #   | test_user    |
        #   | invalid_user |
        # → ["admin", "test_user", "invalid_user"]
        return [row[0] for row in grid if row]

    else:
        raise ValueError(
            f"Unknown as_format='{as_format}'. "
            f"Use: 'list_of_dicts', 'list_of_lists', 'single_dict', 'flat_list'"
        )
